from flask import Blueprint, render_template, request, jsonify, send_file
from datetime import datetime, timedelta
from blueprints.decorators import isloggedin
from .firestore import get_historical_data, get_latest_data
import openpyxl # Ensure openpyxl is imported
from io import BytesIO
from openpyxl.utils import get_column_letter
import pytz # Ensure pytz is imported

# Create a blueprint for history routes
history = Blueprint('history', __name__, url_prefix='/history')

@history.route('/')
@isloggedin
def history_page():
    """Render the history page template"""
    return render_template('history.html')

@history.route('/data')
@isloggedin
def history_data():
    """API endpoint to get historical data"""
    # Get query parameters
    section = request.args.get('section', None)
    days = int(request.args.get('days', 7))
    data_type = request.args.get('type', None)
    
    # Get data from Firestore
    data = get_historical_data(section=section, days=days, data_type=data_type)
    
    return jsonify({
        'success': True,
        'data': data
    })

@history.route('/latest')
@isloggedin
def latest_data():
    """API endpoint to get most recent data"""
    data = get_latest_data()
    
    if data:
        return jsonify({
            'success': True,
            'data': data
        })
    else:
        return jsonify({
            'success': False,
            'message': 'No data available'
        }), 404

@history.route('/export_excel')
@isloggedin
def export_excel():
    """API endpoint to export historical data to Excel"""
    time_range_str = request.args.get('range', '1hour')

    # All defined export ranges are now considered supported with specific handling.
    supported_export_ranges = ['1hour', '1day', '7day', '30day'] 

    if time_range_str not in supported_export_ranges:
        return jsonify({"success": False, "message": "Invalid time range specified."}), 400
    
    days_to_fetch = 0
    if time_range_str == '1hour' or time_range_str == '1day':
        days_to_fetch = 1 # Fetch 1 day of data, then filter for 1hr or use full day for 1day
    elif time_range_str == '7day':
        days_to_fetch = 7
    elif time_range_str == '30day':
        days_to_fetch = 30
    
    raw_data = get_historical_data(days=days_to_fetch)
    processed_data = []
    now_wib = datetime.now(pytz.timezone('Asia/Jakarta'))

    if not raw_data:
        # This message is now handled by the client-side JS for a pop-up
        return jsonify({"success": False, "message": "No raw data found for the selected period."}), 404

    # --- Data Processing based on time_range_str ---
    if time_range_str == '1hour':
        one_hour_ago = now_wib - timedelta(hours=1)
        for record in raw_data:
            record_ts_str = record.get('timestamp')
            if record_ts_str:
                try:
                    record_ts = datetime.fromisoformat(record_ts_str) # Assumes WIB from get_historical_data
                    if record_ts >= one_hour_ago:
                        processed_data.append(record)
                except ValueError:
                    print(f"Warning: Could not parse timestamp {record_ts_str} for 1-hour filter.")
                    continue
    elif time_range_str == '1day':
        # For 1-day export, we want data from the last 24 hours from "now"
        # get_historical_data(days=1) already fetches roughly the last 24 hours of records.
        # We can further refine this to ensure it's strictly the last 24 hours from now_wib if needed,
        # but for simplicity, using the direct output of get_historical_data(days=1) is often sufficient.
        # The current logic filters based on `now_wib - timedelta(days=1)` which is good.
        one_day_ago = now_wib - timedelta(days=1)
        for record in raw_data:
            record_ts_str = record.get('timestamp')
            if record_ts_str:
                try:
                    record_ts = datetime.fromisoformat(record_ts_str) # Assumes WIB
                    if record_ts >= one_day_ago:
                        processed_data.append(record)
                except ValueError:
                    print(f"Warning: Could not parse timestamp {record_ts_str} for 1-day filter.")
                    continue
    elif time_range_str == '7day':
        # Aggregate data into 10-minute intervals for 7-day export
        aggregated_buckets = {} 
        for record in raw_data:
            record_ts_str = record.get('timestamp')
            if not record_ts_str: continue
            try:
                record_ts_wib = datetime.fromisoformat(record_ts_str)
            except ValueError:
                print(f"Warning: Could not parse timestamp {record_ts_str} for 7-day aggregation.")
                continue

            bucket_minute = (record_ts_wib.minute // 10) * 10
            bucket_start_ts = record_ts_wib.replace(minute=bucket_minute, second=0, microsecond=0)
            bucket_start_ts_iso = bucket_start_ts.isoformat()

            if bucket_start_ts_iso not in aggregated_buckets:
                aggregated_buckets[bucket_start_ts_iso] = {'timestamp': bucket_start_ts_iso, 'data': {}}
            
            current_bucket_data = aggregated_buckets[bucket_start_ts_iso]['data']
            for section_name, sensors in record.get('data', {}).items():
                if section_name not in current_bucket_data: current_bucket_data[section_name] = {}
                for sensor_type, values in sensors.items():
                    if not isinstance(values, dict) or 'avg' not in values or 'count' not in values: continue
                    if sensor_type not in current_bucket_data[section_name]:
                        current_bucket_data[section_name][sensor_type] = {
                            'sum_weighted_avg': 0.0, 'total_count': 0,
                            'min_val': float('inf'), 'max_val': float('-inf')
                        }
                    agg_stats = current_bucket_data[section_name][sensor_type]
                    original_avg = values.get('avg', 0.0)
                    original_count = values.get('count', 0)
                    original_min = values.get('min') 
                    original_max = values.get('max')

                    if original_count > 0:
                        agg_stats['sum_weighted_avg'] += original_avg * original_count
                        agg_stats['total_count'] += original_count
                    if original_min is not None and original_min < agg_stats['min_val']:
                        agg_stats['min_val'] = original_min
                    if original_max is not None and original_max > agg_stats['max_val']:
                        agg_stats['max_val'] = original_max
        
        for bucket_ts_iso, bucket_content in sorted(aggregated_buckets.items()):
            final_record_data = {}
            for section_name, sensors_agg in bucket_content['data'].items():
                final_record_data[section_name] = {}
                for sensor_type, agg_values in sensors_agg.items():
                    final_avg = (agg_values['sum_weighted_avg'] / agg_values['total_count']) if agg_values['total_count'] > 0 else None
                    final_min = agg_values['min_val'] if agg_values['min_val'] != float('inf') else None
                    final_max = agg_values['max_val'] if agg_values['max_val'] != float('-inf') else None
                    final_record_data[section_name][sensor_type] = {
                        'avg': final_avg, 'min': final_min, 'max': final_max, 'count': agg_values['total_count']
                    }
            processed_data.append({'timestamp': bucket_ts_iso, 'data': final_record_data})

    elif time_range_str == '30day':
        # Aggregate data into hourly intervals for 30-day export
        aggregated_buckets = {}
        for record in raw_data:
            record_ts_str = record.get('timestamp')
            if not record_ts_str: continue
            try:
                record_ts_wib = datetime.fromisoformat(record_ts_str)
            except ValueError:
                print(f"Warning: Could not parse timestamp {record_ts_str} for 30-day hourly aggregation.")
                continue

            bucket_start_ts = record_ts_wib.replace(minute=0, second=0, microsecond=0) # Hourly bucket
            bucket_start_ts_iso = bucket_start_ts.isoformat()

            if bucket_start_ts_iso not in aggregated_buckets:
                aggregated_buckets[bucket_start_ts_iso] = {'timestamp': bucket_start_ts_iso, 'data': {}}

            current_bucket_data = aggregated_buckets[bucket_start_ts_iso]['data']
            for section_name, sensors in record.get('data', {}).items():
                if section_name not in current_bucket_data: current_bucket_data[section_name] = {}
                for sensor_type, values in sensors.items():
                    if not isinstance(values, dict) or 'avg' not in values or 'count' not in values: continue
                    if sensor_type not in current_bucket_data[section_name]:
                        current_bucket_data[section_name][sensor_type] = {
                            'sum_weighted_avg': 0.0, 'total_count': 0,
                            'min_val': float('inf'), 'max_val': float('-inf')
                        }
                    agg_stats = current_bucket_data[section_name][sensor_type]
                    original_avg = values.get('avg', 0.0)
                    original_count = values.get('count', 0)
                    original_min = values.get('min')
                    original_max = values.get('max')

                    if original_count > 0:
                        agg_stats['sum_weighted_avg'] += original_avg * original_count
                        agg_stats['total_count'] += original_count
                    if original_min is not None and original_min < agg_stats['min_val']:
                        agg_stats['min_val'] = original_min
                    if original_max is not None and original_max > agg_stats['max_val']:
                        agg_stats['max_val'] = original_max
        
        for bucket_ts_iso, bucket_content in sorted(aggregated_buckets.items()):
            final_record_data = {}
            for section_name, sensors_agg in bucket_content['data'].items():
                final_record_data[section_name] = {}
                for sensor_type, agg_values in sensors_agg.items():
                    final_avg = (agg_values['sum_weighted_avg'] / agg_values['total_count']) if agg_values['total_count'] > 0 else None
                    final_min = agg_values['min_val'] if agg_values['min_val'] != float('inf') else None
                    final_max = agg_values['max_val'] if agg_values['max_val'] != float('-inf') else None
                    final_record_data[section_name][sensor_type] = {
                        'avg': final_avg, 'min': final_min, 'max': final_max, 'count': agg_values['total_count']
                    }
            processed_data.append({'timestamp': bucket_ts_iso, 'data': final_record_data})

    if not processed_data:
        # This message is now handled by the client-side JS for a pop-up
        return jsonify({"success": False, "message": "No data to export for the selected range after processing."}), 404

    # --- Excel File Generation ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sensor Data"
    headers = ["Timestamp (WIB)", "Section", "Sensor Type", "Average", "Min", "Max", "Median", "Count"]
    ws.append(headers)

    for record in processed_data:
        timestamp_str = record.get('timestamp')
        try:
            ts_datetime = datetime.fromisoformat(timestamp_str)
            formatted_timestamp = ts_datetime.strftime('%Y-%m-%d %H:%M:%S')
        except (ValueError, TypeError):
            formatted_timestamp = timestamp_str

        sections_data = record.get('data', {})
        for section_name, sensor_types_data in sections_data.items():
            for sensor_type, values in sensor_types_data.items():
                if isinstance(values, dict):
                    
                    # Format section name for display
                    display_section_name = ''
                    if section_name == 'meja_apung':
                        display_section_name = 'Meja Apung'
                    elif section_name == 'peremajaan':
                        display_section_name = 'Peremajaan'
                    elif section_name: # Ensure section_name is not empty or None
                        display_section_name = section_name.capitalize()
                    else:
                        display_section_name = 'Unknown Section'


                    row = [
                        formatted_timestamp,
                        display_section_name,
                        sensor_type.capitalize(),
                        values.get('avg'),
                        values.get('min'),
                        values.get('max'),
                        values.get('median') if time_range_str not in ['7day', '30day'] else None, 
                        values.get('count')
                    ]
                    ws.append(row)
    
    for col_idx, column_cells in enumerate(ws.columns):
        if column_cells: # Ensure column_cells is not empty
            length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[get_column_letter(col_idx + 1)].width = length + 2

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    filename = f"sensor_data_{time_range_str}_{now_wib.strftime('%Y%m%d_%H%M%S')}.xlsx"

    return send_file(
        excel_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )